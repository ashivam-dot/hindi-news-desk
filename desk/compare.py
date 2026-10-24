"""Defence comparison Shorts ("भारत बनाम पाकिस्तान"): two countries' military spending, its share of the economy and
per citizen (SIPRI Military Expenditure Database, read from SIPRI's own spreadsheet and checked against the World
Bank's copy), and their armed forces (IISS, The Military Balance, via the World Bank, CC BY 4.0), over each
country's own freely licensed Commons photos of its forces. No model writes them: every line is a template filled
with a published number, both countries' figures from one year, a SIPRI estimate said to be one, and a figure
SIPRI marks highly uncertain never used. So a comparison cannot say what the data does not.

A pair runs when the day's news is about both countries (a conflict story with OWNERS outlets naming both), or,
when the channel has been quiet for revisit.IDLE, as an evergreen explainer; at most PER_DAY a day and a pair at
most once in PAIR_DAYS days. Owner, 2026-10-24: "some times you can compare the defences of two countries"."""
from __future__ import annotations

import io
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

from . import hindi, media, net
from .spec import Beat, Fact, Short

WB = "https://api.worldbank.org/v2/country/{codes}/indicator/{code}"
WB_PAGE = "https://data.worldbank.org/indicator/{code}?locations={iso2}"
SIPRI_PAGE = "https://www.sipri.org/databases/milex"
SPEND, SHARE, TROOPS, CAPITA = "MS.MIL.XPND.CD", "MS.MIL.XPND.GD.ZS", "MS.MIL.TOTL.P1", "per capita"
SIPRI_SHEETS = {SPEND: "Current US$", SHARE: "Share of GDP", CAPITA: "Per capita"}
# The spreadsheet's legend: "Figures in blue are SIPRI estimates. Figures in red indicate highly uncertain data."
# (Excel's indexed palette: 12 blue, 10 red.)
BLUE, RED = (12, "FF0000FF"), (10, "FFFF0000")
AGREE = 0.005  # the World Bank's copy within 0.5% of SIPRI's own figure, or the comparison is refused
PREFIX, PER_DAY, PAIR_DAYS, OWNERS, FAME = "cmp-", 1, 14, 3, 1
# Off: SIPRI's data (the World Bank's military expenditure series carry "SIPRI terms and conditions", WDI
# metadata License_Type) needs SIPRI's authorization for commercial use, a monetized channel included (2026-10-24).
ENABLED = False
MIN_PHOTOS = 3  # per country: the Short cuts between its forces on land, in the air and at sea
CREDIT = ("तथ्य: SIPRI Military Expenditure Database (रक्षा खर्च); IISS, The Military Balance, World Bank के ज़रिये "
          "(CC BY 4.0) (सशस्त्र बल); तस्वीरें: Wikimedia Commons (लाइसेंस नीचे)")

# iso3: English name, how a headline names it, Hindi, iso2, and the Commons searches for its forces (newest first).
COUNTRIES = {
    "IND": ("India", r"India|Indian", "भारत", "IN",
            ["Indian Army contingent Republic Day parade", "Indian Air Force Rafale", "Indian Navy INS Vikrant",
             "Indian Army T-90 tank"]),
    "PAK": ("Pakistan", r"Pakistan\w*", "पाकिस्तान", "PK",
            ["Pakistan Day Parade", "Pakistan Air Force JF-17 Thunder", "Pakistan Navy frigate PNS",
             "Pakistan Army Al-Khalid tank", "Pakistan Air Force F-16", "Pakistan Army soldiers exercise"]),
    "CHN": ("China", r"China|Chinese|Beijing", "चीन", "CN",
            ["People's Liberation Army honor guard", "PLA Navy ship", "Chengdu J-20", "PLA soldiers exercise"]),
    "USA": ("United States", r"United States|U\.S\.|U S|US|USS|USNS|America\w*|Pentagon", "अमेरिका", "US",
            ["U.S. Army soldiers exercise", "F-35A Lightning II", "USS Gerald R. Ford", "M1 Abrams tank"]),
    "RUS": ("Russia", r"Russia\w*|Kremlin|Moscow", "रूस", "RU",
            ["Moscow Victory Day Parade", "Sukhoi Su-57", "Russian Navy ship", "Russian Army T-90"]),
    "UKR": ("Ukraine", r"Ukrain\w*|Kyiv", "यूक्रेन", "UA",
            ["Armed Forces of Ukraine soldiers training", "Ukrainian Air Force MiG-29", "Ukrainian Navy ship",
             "Ukrainian Ground Forces exercise"]),
    "ISR": ("Israel", r"Israel\w*|IDF", "इज़राइल", "IL",
            ["Israel Defense Forces soldiers", "Israeli Air Force F-35I", "Israeli Navy Sa'ar", "Merkava tank"]),
    "IRN": ("Iran", r"Iran\w*|Tehran", "ईरान", "IR",
            ["Iranian army day parade", "Islamic Republic of Iran Navy ship", "Iranian Air Force F-14",
             "IRGC parade"]),
}
# The pairs viewers here search for, in the order the channel runs them when nothing in the news calls for one.
PAIRS = [("IND", "PAK"), ("IND", "CHN"), ("RUS", "UKR"), ("ISR", "IRN"), ("USA", "CHN"), ("USA", "RUS")]
MIN_YEAR = 2010
NOT_FORCES = re.compile(r"insignia|emblem|flag|logo|badge|patch|seal|coat of arms|medal|roundel|ribbon|uniform "
                        r"pattern|camouflage pattern|diagram|chart|map\b|museum|memorial|grave|cemetery|statue|"
                        r"model|toy|stamp|coin|poster|crest|terroris\w*|Gaza|screenshot|RHIB|"
                        # Not the photo as taken: retouched, or several photos in one.
                        r"altered|edited|retouched|photoshop\w*|composite|collage|montage|clockwise|"
                        # Another country's forces beside them: a joint exercise is not one country's.
                        r"alongside|counterparts|multinational|joint (?:exercise|training)|bilateral", re.I)
# A named person's face is someone the voice never names: no ministers, generals or leaders at a parade.
NAMED = re.compile(r"با حضور|از چپ|از راست|\b(?:From left|from right|L-R|pose[sd]? for|Embassy|delegation|Prime Minister|"
                   r"President|Minister|General|Gen\.|Admiral|Marshal|Chief|Secretary|Sgt\.?|Sergeant|Capt\.|Lt\.|"
                   r"Col\.|Maj\.|Pvt\.|Spc\.|Cpl\.|mentor|Modi|"
                   r"Putin|Xi Jinping|Biden|Trump|Netanyahu|Khamenei|Zelensky\w*|Shoigu|Hegseth|Austin|Sharif|"
                   r"Munir|Rajnath|Singh)(?!\w)")


def _iso3(text: str) -> list[str]:
    return [k for k, v in COUNTRIES.items() if re.search(rf"\b(?:{v[1]})(?!\w)", text or "")]


SIPRI_NAMES = {"USA": "United States of America"}


def sipri_name(iso3: str) -> str:
    return SIPRI_NAMES.get(iso3, COUNTRIES[iso3][0])


def _colour(cell) -> str:
    """"estimate" (blue), "uncertain" (red) or "" for a SIPRI cell."""
    c = cell.font.color if cell.font else None
    if c is None:
        return ""
    got = c.indexed if c.type == "indexed" else (c.rgb if c.type == "rgb" else None)
    return "estimate" if got in BLUE else "uncertain" if got in RED else ""


def sipri_tables(content: bytes) -> dict[str, dict[str, dict[int, tuple[float, str]]]]:
    """{sheet key: {SIPRI country name: {year: (value, "estimate" | "uncertain" | "")}}} from SIPRI's spreadsheet;
    a cell without a number ("...", "xxx") is no figure."""
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(content))
    out = {}
    for key, sheet in SIPRI_SHEETS.items():
        ws = book[sheet]
        head = next(r for r in ws.iter_rows() if r[0].value == "Country")
        years = {c.column: int(float(c.value)) for c in head if re.fullmatch(r"\d{4}(\.0)?", str(c.value or ""))}
        table: dict[str, dict] = {}
        for row in ws.iter_rows(min_row=head[0].row + 1):
            name = str(row[0].value or "").strip()
            if name:
                table[name] = {years[c.column]: (float(c.value), _colour(c)) for c in row
                               if c.column in years and isinstance(c.value, (int, float))}
        out[key] = table
    return out


def sipri_file(get=net.get) -> bytes:
    """SIPRI's current spreadsheet, found from its database page (the file name carries the version)."""
    page = get(SIPRI_PAGE, timeout=60).decode("utf-8", "replace")
    m = re.search(r'href="((?:https:)?//www\.sipri\.org/[^"]*Milex-data[^"]*\.xlsx)"', page)
    if not m:
        raise ValueError("compare: SIPRI's database page links no spreadsheet")
    return get(("https:" + m.group(1)) if m.group(1).startswith("//") else m.group(1), timeout=120)


def _wb(pair, code, get) -> dict[str, dict[int, float]]:
    d = get(WB.format(codes=";".join(pair), code=code), params={"format": "json", "mrv": 8, "per_page": 100},
            timeout=60)
    out: dict[str, dict[int, float]] = {k: {} for k in pair}
    for row in (d[1] if isinstance(d, list) and len(d) > 1 else None) or []:
        if row.get("countryiso3code") in out and row.get("value") is not None:
            out[row["countryiso3code"]][int(row["date"])] = float(row["value"])
    return out


def figures(pair, get=net.get_json, sipri: dict | None = None) -> dict:
    """{"year", "troops_year", iso3: {SPEND, SHARE, CAPITA, TROOPS, "estimate"}}: the latest year SIPRI has all three
    spending figures for both countries, none highly uncertain, each spending figure within AGREE of the World Bank's
    copy; the latest year IISS (World Bank) has both countries' armed forces. Refuses rather than mixing years."""
    sipri = sipri if sipri is not None else sipri_tables(sipri_file())
    rows = {key: {k: sipri[key].get(sipri_name(k), {}) for k in pair} for key in SIPRI_SHEETS}
    usable = [y for y in sorted(set.intersection(*(set(rows[key][k]) for key in rows for k in pair)), reverse=True)
              if all(rows[key][k][y][1] != "uncertain" for key in rows for k in pair)]
    if not usable:
        raise ValueError(f"compare: SIPRI has no common year of firm figures for {' and '.join(pair)}")
    year = usable[0]
    wb = {SPEND: _wb(pair, SPEND, get), SHARE: _wb(pair, SHARE, get)}
    for k in pair:
        ours, theirs = wb[SPEND][k].get(year), rows[SPEND][k][year][0] * 1e6  # SIPRI: US$ m.
        if ours is None or abs(ours - theirs) > AGREE * theirs:
            raise ValueError(f"compare: {k} {year} spending: World Bank {ours} against SIPRI {theirs:.0f}")
        ours, theirs = wb[SHARE][k].get(year), rows[SHARE][k][year][0] * 100  # SIPRI: a fraction
        if ours is None or abs(ours - theirs) > max(0.05, AGREE * theirs):
            raise ValueError(f"compare: {k} {year} share of GDP: World Bank {ours} against SIPRI {theirs:.2f}")
    troops = _wb(pair, TROOPS, get)
    both = sorted(set(troops[pair[0]]) & set(troops[pair[1]]), reverse=True)
    if not both:
        raise ValueError(f"compare: IISS (World Bank) has no common year of armed forces for {' and '.join(pair)}")
    out: dict = {"year": year, "troops_year": both[0]}
    for k in pair:
        out[k] = {SPEND: rows[SPEND][k][year][0] * 1e6, SHARE: rows[SHARE][k][year][0] * 100,
                  CAPITA: rows[CAPITA][k][year][0], TROOPS: troops[k][both[0]],
                  "estimate": any(rows[key][k][year][1] == "estimate" for key in rows)}
    return out


def _one(x: float) -> str:
    """One decimal, without a trailing .0 ("7.8", "12")."""
    s = f"{x:.1f}"
    return s[:-2] if s.endswith(".0") else s


def money(usd: float) -> str:
    """"92.1 अरब डॉलर", "27.8 करोड़ डॉलर"."""
    return f"{_one(usd / 1e9)} अरब डॉलर" if usd >= 1e9 else f"{_one(usd / 1e7)} करोड़ डॉलर"


def people(n: float) -> str:
    """"30.7 लाख", "94 हज़ार"."""
    return f"{_one(n / 1e5)} लाख" if n >= 1e5 else f"{round(n / 1e3)} हज़ार"


def times(a: float, b: float) -> str:
    return _one(a / b)


def script(pair, data) -> dict:
    """The title, the beats ({say, show: [(iso3, slot)], card}) and the facts ({key, phrase, code, iso3s, what})."""
    a, b = sorted(pair, key=lambda k: -data[k][SPEND])  # the bigger spender leads
    A, B = COUNTRIES[a][2], COUNTRIES[b][2]
    ys, yt = data["year"], data["troops_year"]
    sa, sb = data[a][SPEND], data[b][SPEND]
    ta, tb = data[a][TROOPS], data[b][TROOPS]
    pa, pb = round(data[a][CAPITA]), round(data[b][CAPITA])
    r = times(sa, sb)
    hi_share, lo_share = (a, b) if data[a][SHARE] >= data[b][SHARE] else (b, a)
    H, L = COUNTRIES[hi_share][2], COUNTRIES[lo_share][2]
    gh, gl = _one(data[hi_share][SHARE]), _one(data[lo_share][SHARE])
    lead_in = "और जीडीपी के हिस्से में भी" if hi_share == a else "पर जीडीपी के हिस्से में"
    guessed = [COUNTRIES[k][2] for k in (a, b) if data[k]["estimate"]]
    # SIPRI's own estimate, said as one (SIPRI estimates China's, Russia's and Ukraine's 2025 spending).
    estimate = f" {' और '.join(guessed)} के आंकड़े सिपरी का अनुमान हैं।" if guessed else ""
    mark = " · अनुमान" if data[a]["estimate"] else ""  # the card shows the lead's figure: its own mark only
    # About 65 words, about 32 s at the narrator's 115 a minute (dry run 40103609151: 70 words ran 38 s).
    beats = [
        {"say": f"सिपरी के मुताबिक, रक्षा पर {A} का खर्च {B} से {r} गुना है।", "show": [(a, 0), (b, 0)], "cut": B,
         "card": {"big": f"{r} गुना", "small": f"रक्षा खर्च · सिपरी · {ys}"}},
        {"say": f"{ys} में {A} ने {money(sa)} और {B} ने {money(sb)} खर्च किए।{estimate}",
         "show": [(a, 1), (b, 1)], "cut": B, "card": {"big": money(sa), "small": f"{A} · {ys}{mark}"}},
        {"say": f"{lead_in} {H} आगे है: {gh} प्रतिशत, {L} {gl} प्रतिशत।",
         "show": [(hi_share, 2), (lo_share, 2)], "cut": L,
         "card": {"big": f"{gh}% बनाम {gl}%", "small": f"जीडीपी में हिस्सा · {ys}"}},
        {"say": f"{yt} में {A} के सशस्त्र बल {people(ta)} और {B} के {people(tb)} थे।",
         "show": [(a, 3), (b, 3)], "cut": B,
         "card": {"big": f"{people(ta)} बनाम {people(tb)}", "small": f"सशस्त्र बल, अर्धसैनिक समेत · आईआईएसएस · {yt}"}},
        {"say": f"हर नागरिक पर {A} {pa} डॉलर और {B} {pb} डॉलर खर्च करता है।",
         "show": [(a, 4), (b, 4)], "cut": B,
         "card": {"big": f"{pa} बनाम {pb} डॉलर", "small": f"प्रति नागरिक · सिपरी · {ys}"}},
    ]
    sipri = lambda k, what: {"source": SIPRI_PAGE, "code": SPEND, "iso3s": k,  # noqa: E731
                             "what": f"SIPRI Military Expenditure Database, {ys}: {what}" +
                                     (" (SIPRI estimate)" if any(data[x]["estimate"] for x in k) else "")}
    facts = [
        {"key": "spend-ratio", "phrase": f"{ys} {r} गुना",
         **sipri((a, b), f"{sa:,.0f} / {sb:,.0f} current US$ = {sa / sb:.2f}")},
        *({"key": f"spend-{k}", "phrase": f"{ys} {money(data[k][SPEND])}",
           **sipri((k,), f"{data[k][SPEND]:,.0f} current US$")} for k in (a, b)),
        *({"key": f"share-{k}", "phrase": f"{ys} {_one(data[k][SHARE])} प्रतिशत",
           **sipri((k,), f"{data[k][SHARE]:.2f}% of GDP")} for k in (a, b)),
        *({"key": f"per-person-{k}", "phrase": f"{ys} {round(data[k][CAPITA])} डॉलर",
           **sipri((k,), f"{data[k][CAPITA]:.2f} current US$ per capita")} for k in (a, b)),
        *({"key": f"troops-{k}", "phrase": f"{yt} {people(data[k][TROOPS])}", "code": TROOPS, "iso3s": (k,),
           "source": WB_PAGE.format(code=TROOPS, iso2=COUNTRIES[k][3]),
           "what": f"IISS, The Military Balance (World Bank WDI {TROOPS}), {yt}: {data[k][TROOPS]:,.0f} armed forces "
                   f"personnel, paramilitary included"} for k in (a, b)),
    ]
    title = f"{A} बनाम {B}: रक्षा खर्च में {r} गुना का फ़र्क"
    return {"lead": a, "other": b, "title": title, "beats": beats, "facts": facts, "year": ys}


def own(p, iso3: str) -> bool:
    """The photo is of this country's forces by its own words: its file or description names the country and no
    other listed one. A "Pakistan Army parade" search returned an Indian Army regiment's bicentenary (2026-10-24)."""
    text = f"{p.file} {p.description or ''} {' '.join(p.categories or [])}"
    return iso3 in _iso3(f"{p.file} {p.description or ''}") and set(_iso3(text)) <= {iso3}


def forces_photos(iso3: str, folder: Path, taken: set[str], recent: set[str] = frozenset(),
                  want: int = 5) -> list[dict]:
    """Up to `want` free Commons photos of a country's forces, one search after another, MIN_YEAR or newer, none of
    an emblem or a named person, preferring photos the channel has not shown lately."""
    pools: list[list] = []
    for q in COUNTRIES[iso3][4]:
        try:
            files = media._search_files(q, 25)
        except Exception:
            continue
        picked = [p for p in media.picks(files, 12)
                  if (media.year_of(p) or 0) >= MIN_YEAR and own(p, iso3)
                  and not NOT_FORCES.search(f"{p.file} {p.description}") and not NAMED.search(p.description or "")]
        pools.append(sorted(picked, key=lambda p: p.file in recent))
    # One from each search in turn (a parade, a jet, a ship, a tank), not two looks at one jet.
    found: list = []
    for rank in range(max(map(len, pools), default=0)):
        for pool in pools:
            if len(found) < want and rank < len(pool) and pool[rank].file not in taken:
                found.append(pool[rank])
                taken.add(pool[rank].file)
    return media.as_dicts(media.download(found, folder))


def build(cand: dict, get=net.get_json, photos_for=forces_photos, sipri: dict | None = None) -> Short:
    pair = tuple(cand["pair"])
    data = figures(pair, get, sipri)
    s = script(pair, data)
    from . import writer

    recent, taken = writer.recent_photos(), set()
    folder = writer.WORK / cand["id"] / "photos"
    photos, own = [], {}
    for k in (s["lead"], s["other"]):
        mine = photos_for(k, folder / k, taken, recent)
        if len(mine) < MIN_PHOTOS:
            raise ValueError(f"compare: only {len(mine)} free photos of {COUNTRIES[k][0]}'s forces")
        own[k] = list(range(len(photos), len(photos) + len(mine)))
        for p in mine:
            # tag: the country stamped on every view of it (compose.places_on); no "who", whose one-off name strap
            # would repeat it.
            p.update(role="thing", tag=COUNTRIES[k][2], depicts=[hindi.norm(COUNTRIES[k][2])])
        photos += mine
    cast = [{"en": COUNTRIES[k][0], "hi": COUNTRIES[k][2], "kind": "thing", "photos": own[k]} for k in own]
    beats = []
    for b in s["beats"]:
        shown = [own[k][slot % len(own[k])] for k, slot in b["show"]]
        cut = b.get("cut") or ""
        at = b["say"].rfind(cut) if cut else -1
        beats.append(Beat(b["say"], "photo:" + "/".join(map(str, shown)), b["card"], cut_word=cut, cut_at=at))
    facts = []
    for f in s["facts"]:
        facts.append(Fact(f["key"], f["phrase"], f["source"], [v for v, _ in hindi.numbers(f["phrase"])], f["what"]))
    sources = list(dict.fromkeys(f.source for f in facts))
    A, B = COUNTRIES[s["lead"]][2], COUNTRIES[s["other"]][2]
    en = f"{COUNTRIES[s['lead']][0]} vs {COUNTRIES[s['other']][0]}"
    credits = [CREDIT] + sorted({p["credit"] + " " + p["page"] for p in photos})
    region = "india" if "IND" in pair else "world"
    scene = {"template": "full", "mood": "news", "voice": "sulafat", "region": region, "photos": photos,
             "name_hi": "", "subject_person": False, "names": [], "cast": cast, "category": "defence",
             "category_hi": "रक्षा तुलना", "lane": "compare", "pair": list(pair), "min_photos": 2 * MIN_PHOTOS,
             "hook": "", "writer": {"models": ["template"], "rejected": [], "titles": [s["title"]]}}
    tags = [f"{en} military", f"{en} defence budget", f"{COUNTRIES[s['lead']][0]} military spending",
            f"{COUNTRIES[s['other']][0]} military spending", "defence comparison", "hindi news", "news in hindi"]
    slug = lambda k: re.sub(r"[^a-z]", "", COUNTRIES[k][0].lower())  # noqa: E731
    hashtags = [f"{slug(s['lead'])}vs{slug(s['other'])}", "defence", "military", "hindinews", "shorts",
                slug(s["lead"]), slug(s["other"]), "worldnews" if region == "world" else "indianews", "रक्षा"]
    summary = (f"{A} और {B} की सेनाओं की तुलना: रक्षा खर्च, जीडीपी में हिस्सा, सशस्त्र बल और प्रति नागरिक खर्च। "
               f"रक्षा खर्च, जीडीपी में हिस्सा और प्रति नागरिक खर्च ({s['year']}): SIPRI Military Expenditure Database; "
               f"सशस्त्र बल ({data['troops_year']}, अर्धसैनिक बल समेत): IISS, The Military Balance (World Bank)।"
               + (" SIPRI के अनुमान वाले आंकड़े अनुमान ही बताए गए हैं।"
                  if any(data[k]["estimate"] for k in pair) else ""))
    return Short(id=cand["id"], kind="compare", title=s["title"], beats=beats, facts=facts, sources=sources,
                 loss=False, scene=scene, event_time=f"{s['year']}", tags=tags, summary=summary, credits=credits,
                 hashtags=hashtags)


def _day(at: datetime) -> str:
    return (at + timedelta(hours=5, minutes=30)).strftime("%Y%m%d")


def _pair_of(sid: str) -> tuple[str, str] | None:
    m = re.match(rf"{PREFIX}\d{{8}}-([A-Z]{{3}})-([A-Z]{{3}})", sid or "")
    return (m.group(1), m.group(2)) if m else None


def news_pairs(stories: list[dict]) -> list[tuple[tuple[str, str], dict]]:
    """The listed pairs today's well-reported stories name both countries of: (pair, story), best covered first."""
    out = []
    for c in stories:
        if len(c.get("owners") or []) < OWNERS:
            continue
        text = " ".join([c.get("title") or ""] + [h if isinstance(h, str) else h.get("title", "")
                                                   for h in c.get("headlines") or []])
        named = set(_iso3(text))
        out += [(p, c) for p in PAIRS if set(p) <= named]
    return sorted(out, key=lambda pc: -len(pc[1].get("owners") or []))


def candidate(pair, at: datetime, why: str, story: dict | None = None) -> dict:
    a, b = pair
    title = f"{COUNTRIES[a][0]} and {COUNTRIES[b][0]} military comparison"
    return {"id": f"{PREFIX}{_day(at)}-{a}-{b}", "kind": "compare", "compare": True, "pair": [a, b], "title": title,
            "name": title, "category": "defence", "region": "india" if "IND" in pair else "world",
            "date": at.date().isoformat(), "news_at": at.isoformat(), "fame": FAME, "owners": [],
            "headlines": [], "top": None, "why": why, "story": (story or {}).get("id")}


def candidates(stories: list[dict], rows: list[dict], at: datetime) -> tuple[list[dict], list[str]]:
    """Today's comparison, if one is due: a pair in the news first, else the next evergreen pair while the channel is
    quiet (nothing published for revisit.IDLE); none when PER_DAY ran today, none for a pair told within
    PAIR_DAYS."""
    from . import revisit

    if not ENABLED:
        return [], []
    day = _day(at)
    mine = [r for r in rows if (r.get("id") or "").startswith(PREFIX)]
    if sum(_day(_at(r) or at) == day for r in mine) >= PER_DAY:
        return [], ["compare: one ran today"]
    since = at - timedelta(days=PAIR_DAYS)
    told = {frozenset(p) for r in mine if (p := _pair_of(r["id"])) and (_at(r) or at) >= since}
    for pair, story in news_pairs(stories):
        if frozenset(pair) not in told:
            return [candidate(pair, at, "in the news", story)], []
    last_any = max((t for r in rows if (t := _at(r))), default=None)
    if last_any is None or at - last_any < revisit.IDLE:
        return [], []
    last = {frozenset(p): _at(r) for r in mine if (p := _pair_of(r["id"]))}
    fresh = [p for p in PAIRS if frozenset(p) not in told]
    if not fresh:
        return [], ["compare: every pair ran in the last 14 days"]
    fresh.sort(key=lambda p: last.get(frozenset(p)) or datetime.min.replace(tzinfo=timezone.utc))
    return [candidate(fresh[0], at, "the channel is quiet")], []


def _at(row: dict) -> datetime | None:
    for key in ("published_at", "at", "made_at"):
        if v := row.get(key):
            try:
                return datetime.fromisoformat(str(v).replace("Z", "+00:00"))
            except ValueError:
                continue
    return None


def from_target(target: str, at: datetime | None = None) -> dict:
    """"compare:IND-PAK" (desk make) as a candidate."""
    a, b = target.split(":", 1)[1].upper().split("-")
    if a not in COUNTRIES or b not in COUNTRIES:
        raise ValueError(f"compare: no data list for {a} or {b} (have {', '.join(COUNTRIES)})")
    return candidate((a, b), at or datetime.now(timezone.utc), "asked for")
