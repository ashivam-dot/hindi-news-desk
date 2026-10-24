from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest

from desk import compare, gate, hindi, writer

AT = datetime(2026, 10, 24, 11, 0, tzinfo=timezone.utc)
# The World Bank's copy (WDI): spending 2024-2025, personnel (IISS) 2019-2020.
WB_ROWS = {
    compare.SPEND: {"IND": {2025: 92057013588.16, 2024: 85597083119.11},
                    "PAK": {2025: 11871380866.11, 2024: 10303476993.98},
                    "CHN": {2025: 335524155624.44, 2024: 311900350600.82}},
    compare.SHARE: {"IND": {2025: 2.3, 2024: 2.22}, "PAK": {2025: 2.91, 2024: 2.73}, "CHN": {2025: 1.72, 2024: 1.66}},
    compare.TROOPS: {"IND": {2020: 3068000, 2019: 3045000}, "PAK": {2020: 943000, 2019: 943000},
                     "CHN": {2020: 2535000}},
}
# SIPRI's own spreadsheet, as sipri_tables reads it: US$ m., a fraction of GDP, US$ per capita; China's in blue.
SIPRI = {
    compare.SPEND: {"India": {2025: (92057.0135881601, ""), 2024: (85597.08, "")},
                    "Pakistan": {2025: (11871.380866105368, ""), 2024: (10303.48, "")},
                    "China": {2025: (335524.15562444, "estimate"), 2024: (311900.35, "estimate")}},
    compare.SHARE: {"India": {2025: (0.022951791793515173, ""), 2024: (0.0222, "")},
                    "Pakistan": {2025: (0.029136125148539926, ""), 2024: (0.0273, "")},
                    "China": {2025: (0.01724723316575711, "estimate"), 2024: (0.0166, "estimate")}},
    compare.CAPITA: {"India": {2025: (63.16370351942166, ""), 2024: (59.0, "")},
                     "Pakistan": {2025: (46.88488347272678, ""), 2024: (41.0, "")},
                     "China": {2025: (236.66233706007665, "estimate"), 2024: (220.0, "estimate")}},
}


def fake_get(url, params=None, timeout=None, rows=WB_ROWS):
    code = url.rsplit("/", 1)[1]
    codes = url.split("/country/")[1].split("/")[0].split(";")
    return [{}, [{"countryiso3code": k, "date": str(y), "value": v} for k in codes
                 for y, v in rows[code].get(k, {}).items()]]


def fake_photos(iso3, folder, taken, recent=frozenset(), want=5):
    en = compare.COUNTRIES[iso3][0]
    return [{"file": f"File:{en} forces {i}.jpg", "page": f"https://commons.wikimedia.org/wiki/File:{en}_{i}.jpg",
             "license": "CC BY 4.0", "credit": f"फ़ोटो: A / CC BY 4.0 / Wikimedia Commons", "path": "", "year": 2023,
             "description": f"{en} Army parade"} for i in range(want)]


def test_a_comparison_is_the_published_numbers_and_passes_every_gate(monkeypatch, tmp_path):
    monkeypatch.setattr(writer, "WORK", tmp_path)
    monkeypatch.setattr(writer, "recent_photos", lambda: set())
    short = compare.build(compare.from_target("compare:PAK-IND", AT), get=fake_get, photos_for=fake_photos,
                          sipri=SIPRI)
    assert short.title == "भारत बनाम पाकिस्तान: रक्षा खर्च में 7.8 गुना का फ़र्क"  # the bigger spender leads
    said = " ".join(b.text for b in short.beats)
    for words in ("सिपरी के मुताबिक", "7.8 गुना", "92.1 अरब डॉलर", "11.9 अरब डॉलर", "पाकिस्तान आगे है: 2.9 प्रतिशत",
                  "2020 में", "30.7 लाख", "9.4 लाख", "63 डॉलर", "47 डॉलर"):  # SIPRI's per-capita, not ours
        assert words in said
    assert "अनुमान" not in said  # India's and Pakistan's 2025 figures are reported, not SIPRI estimates
    assert all(p["tag"] in ("भारत", "पाकिस्तान") and "who" not in p for p in short.scene["photos"])
    backed = {n for f in short.facts for n in f.numbers}
    assert {v for v, _ in hindi.numbers(said + " " + short.title)} <= backed
    by = {f.key: f for f in short.facts}
    assert by["spend-IND"].source == compare.SIPRI_PAGE and "IISS" in by["troops-IND"].evidence
    assert "IISS" in short.credits[0] and "SIPRI" in short.credits[0] and "IISS" in short.summary
    assert gate.check(short) == [] and gate.said_is_shown(short) == []


def test_a_sipri_estimate_is_said_to_be_one(monkeypatch, tmp_path):
    monkeypatch.setattr(writer, "WORK", tmp_path)
    monkeypatch.setattr(writer, "recent_photos", lambda: set())
    short = compare.build(compare.from_target("compare:IND-CHN", AT), get=fake_get, photos_for=fake_photos,
                          sipri=SIPRI)
    china = next(c["photos"] for c in short.scene["cast"] if c["en"] == "China")
    assert short.beats[2].text == "चीन के आंकड़े सिपरी का अनुमान हैं।"
    assert int(short.beats[2].shot.split(":")[1]) in china  # over China's own photo, not India's
    assert gate.check(short) == [] and gate.said_is_shown(short) == []
    assert "अनुमान" in short.beats[1].card["small"] and "(SIPRI estimate)" in {f.key: f for f in short.facts}[
        "spend-CHN"].evidence
    sipri_us = {k: t | {"United States of America": {2025: (954387.0, ""), 2024: (1004901.0, "")}}
                for k, t in SIPRI.items()}
    sipri_us[compare.SHARE]["United States of America"] = {2025: (0.031173, ""), 2024: (0.0343, "")}
    sipri_us[compare.CAPITA]["United States of America"] = {2025: (2755.37, ""), 2024: (2950.0, "")}
    rows = {k: dict(v) for k, v in WB_ROWS.items()}
    rows[compare.SPEND] = rows[compare.SPEND] | {"USA": {2025: 954387000000, 2024: 1004901000000}}
    rows[compare.SHARE] = rows[compare.SHARE] | {"USA": {2025: 3.12, 2024: 3.43}}
    rows[compare.TROOPS] = rows[compare.TROOPS] | {"USA": {2020: 1395000}}
    s = compare.script(("USA", "CHN"), compare.figures(("USA", "CHN"), lambda *a, **k: fake_get(*a, rows=rows, **k),
                                                       sipri_us))
    assert "अनुमान" in s["beats"][2]["say"] and "अनुमान" not in s["beats"][1]["card"]["small"]  # the US card


def test_figures_never_mix_years_or_use_uncertain_or_unmatched_data():
    data = compare.figures(("IND", "PAK"), fake_get, SIPRI)
    assert (data["year"], data["troops_year"]) == (2025, 2020)
    shaky = {k: {n: dict(v) for n, v in t.items()} for k, t in SIPRI.items()}
    shaky[compare.SPEND]["Pakistan"][2025] = (11871.38, "uncertain")
    assert compare.figures(("IND", "PAK"), fake_get, shaky)["year"] == 2024  # both countries' last firm year
    off = {k: {c: dict(v) for c, v in t.items()} for k, t in WB_ROWS.items()}
    off[compare.SPEND]["PAK"][2025] = 12_500_000_000  # 5% off SIPRI's own figure
    with pytest.raises(ValueError, match="World Bank .* against SIPRI"):
        compare.figures(("IND", "PAK"), lambda *a, **k: fake_get(*a, rows=off, **k), SIPRI)
    lone = {k: {c: dict(v) for c, v in t.items()} for k, t in WB_ROWS.items()}
    lone[compare.TROOPS]["PAK"] = {2019: 943000}
    assert compare.figures(("IND", "PAK"), lambda *a, **k: fake_get(*a, rows=lone, **k), SIPRI)[
        "troops_year"] == 2019


def test_sipri_tables_read_the_estimate_and_uncertain_colours(tmp_path):
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.styles.colors import Color

    book = openpyxl.Workbook()
    book.remove(book.active)
    for sheet in compare.SIPRI_SHEETS.values():
        ws = book.create_sheet(sheet)
        ws.append(["Figures in blue are SIPRI estimates."])
        ws.append(["Country", "Notes", 2024, 2025])
        ws.append(["India", None, 85597.08, 92057.01])
        ws.append(["China", None, 311900.35, 335524.16])
        ws.append(["Korea, North", None, "...", 4000.0])
        ws["D4"].font = Font(color=Color(indexed=12))
        ws["D5"].font = Font(color=Color(indexed=10))
    path = tmp_path / "milex.xlsx"
    book.save(path)
    t = compare.sipri_tables(path.read_bytes())[compare.SPEND]
    assert t["India"] == {2024: (85597.08, ""), 2025: (92057.01, "")}
    assert t["China"][2025] == (335524.16, "estimate") and t["Korea, North"] == {2025: (4000.0, "uncertain")}


def test_a_photo_must_name_its_own_country_and_no_other():
    def p(file, description="", categories=()):
        return SimpleNamespace(file=file, description=description, categories=list(categories))

    assert compare.own(p("File:Pakistan Air Force JF-17 Thunder Block III.jpg"), "PAK")
    # A "Pakistan Army parade" hit: an Indian Army regiment's bicentenary.
    assert not compare.own(p("File:Bicentenary celebrations of The Poona Horse3.jpg",
                             "DECORATED REGIMENT OF INDIAN ARMY CELEBRATED THE BICENTENARY", ["Pakistan"]), "PAK")
    assert not compare.own(p("File:Russian Navy ship in Vladivostok.jpg",
                             "VLADIVOSTOK, Russia (Sept. 22, 2012) Sailors aboard USS Vandegrift"), "RUS")
    assert not compare.own(p("File:Ukrainian MiG-29s.jpg", "A U.S. B-52H Stratofortress and Ukrainian MiG-29s"), "UKR")
    # A city's parade is not the country's forces: Belarus's contingent at Moscow's 2010 parade.
    assert not compare.own(p("File:2010 Moscow Victory Day Parade-9.jpeg",
                             "Военный парад, посвящённый 65-летию Победы"), "RUS")
    assert compare.own(p("File:Sukhoi Su-57 RF-81775 Army-2022.jpg", "Russian Air Force Sukhoi Su-57"), "RUS")
    # A named person is someone the voice never names.
    assert compare.NAMED.search("Christmas Eve photo of Andrii Pilshchykov posted by the Ukrainian Air Force")
    assert not compare.NAMED.search("Ukrainian Air Force MiG-29 in the colours of the Ukrainian Falcons")


def rows_at(*items):
    return [{"id": sid, "published_at": at.isoformat()} for sid, at in items]


def test_a_quiet_channel_gets_the_next_evergreen_pair_once_a_day(monkeypatch):
    monkeypatch.setattr(compare, "ENABLED", True)
    quiet = rows_at(("kh-20261024-Q1", AT - timedelta(hours=3)))
    (c,), _ = compare.candidates([], quiet, AT)
    assert c["id"] == "cmp-20261024-IND-PAK" and c["compare"] and c["why"] == "the channel is quiet"
    busy = rows_at(("kh-20261024-Q1", AT - timedelta(minutes=30)))
    assert compare.candidates([], busy, AT)[0] == []
    ran = quiet + rows_at(("cmp-20261024-IND-PAK", AT - timedelta(hours=4)))
    assert compare.candidates([], ran, AT) == ([], ["compare: one ran today"])
    yesterday = quiet + rows_at(("cmp-20261023-IND-PAK", AT - timedelta(days=1)))
    (c,), _ = compare.candidates([], yesterday, AT)
    assert c["pair"] == ["IND", "CHN"]  # a pair waits PAIR_DAYS


def test_a_pair_in_the_news_runs_even_while_the_channel_is_busy(monkeypatch):
    monkeypatch.setattr(compare, "ENABLED", True)
    busy = rows_at(("kh-20261024-Q1", AT - timedelta(minutes=30)))
    story = {"id": "nw-x", "title": "Israel strikes Iran's air defences", "owners": ["a", "b", "c"],
             "headlines": ["Iran vows response after Israeli strikes"]}
    (c,), _ = compare.candidates([story], busy, AT)
    assert c["pair"] == ["ISR", "IRN"] and c["why"] == "in the news" and c["story"] == "nw-x"
    assert compare.candidates([story | {"owners": ["a"]}], busy, AT)[0] == []


@pytest.mark.parametrize("usd,said", [(92057013588, "92.1 अरब डॉलर"), (278000000, "27.8 करोड़ डॉलर")])
def test_money_reads_in_arab_or_crore(usd, said):
    assert compare.money(usd) == said
